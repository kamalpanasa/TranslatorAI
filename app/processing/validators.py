import os
import fitz  # PyMuPDF
import pandas as pd
from PIL import Image
from docx import Document
from openpyxl import load_workbook
from app.utils.exceptions import FileProcessingError
from app.utils.logger import get_logger

logger = get_logger("app.processing.validators")


def validate_pdf(filepath: str) -> None:
    """
    Validates that a PDF is not encrypted, corrupted, or empty.
    """
    try:
        with fitz.open(filepath) as doc:
            if doc.is_encrypted:
                raise FileProcessingError("Encrypted PDFs are not supported. Please remove password protection.")
            if doc.page_count == 0:
                raise FileProcessingError("The uploaded PDF has 0 pages.")
    except FileProcessingError as fpe:
        raise fpe
    except Exception as e:
        logger.error(f"PDF validation failed: {str(e)}")
        raise FileProcessingError("The uploaded PDF is corrupted or invalid.")


def validate_image(filepath: str) -> None:
    """
    Validates that an image is readable by Pillow.
    """
    try:
        with Image.open(filepath) as img:
            img.verify()
    except Exception as e:
        logger.error(f"Image validation failed: {str(e)}")
        raise FileProcessingError("The uploaded image is corrupted or invalid.")


def validate_docx(filepath: str) -> None:
    """
    Validates that a DOCX is readable.
    """
    try:
        Document(filepath)
    except Exception as e:
        logger.error(f"DOCX validation failed: {str(e)}")
        raise FileProcessingError("The uploaded Word document is corrupted or invalid.")


def validate_excel(filepath: str) -> None:
    """
    Validates that an Excel workbook is readable.
    """
    try:
        wb = load_workbook(filepath, read_only=True)
        wb.close()
    except Exception as e:
        logger.error(f"Excel validation failed: {str(e)}")
        raise FileProcessingError("The uploaded Excel file is corrupted or invalid.")


def validate_csv(filepath: str) -> None:
    """
    Validates that a CSV is readable.
    """
    try:
        # Just try to read the first few rows to verify separator structure
        pd.read_csv(filepath, nrows=5)
    except Exception as e:
        logger.error(f"CSV validation failed: {str(e)}")
        raise FileProcessingError("The uploaded CSV file is corrupted or invalid.")
